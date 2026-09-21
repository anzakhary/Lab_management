from __future__ import annotations

import json
import logging
import os
import smtplib
from datetime import date, datetime
from email.mime.multipart import MIMEMultipart
from email.mime.text import MIMEText
from email.utils import formatdate
from pathlib import Path

from flask import Flask, redirect, render_template, request, send_from_directory, session, url_for
from openpyxl import load_workbook
from werkzeug.exceptions import HTTPException
from werkzeug.security import check_password_hash, generate_password_hash

app = Flask(__name__)
app.config.update(
    SECRET_KEY=os.environ.get("LAB_SECRET_KEY") or os.urandom(32).hex(),
    SESSION_COOKIE_HTTPONLY=True,
    SESSION_COOKIE_SAMESITE="Lax",
    SESSION_COOKIE_SECURE=os.environ.get("LAB_SESSION_COOKIE_SECURE", "false").lower() == "true",
)

ADMIN_USERNAME = (os.environ.get("LAB_ADMIN_USERNAME") or "admin").strip()
ADMIN_PASSWORD = (os.environ.get("LAB_ADMIN_PASSWORD") or "").strip()
ADMIN_PASSWORD_HASH = generate_password_hash(ADMIN_PASSWORD) if ADMIN_PASSWORD else None

ROOT = Path(__file__).resolve().parent
WORKBOOK_PATH = ROOT / "Lab_Inventory_Tracker (2).xlsx"
REQUESTS_PATH = ROOT / "borrow_requests.json"
BORROWED_PATH = ROOT / "borrowed_parts.json"
LOG_PATH = ROOT / "lab_inventory.log"
MAIL_SERVER = os.environ.get("MAIL_SERVER", "smtp.gmail.com")
MAIL_PORT = int(os.environ.get("MAIL_PORT", "587"))
MAIL_USERNAME = os.environ.get("MAIL_USERNAME", "").strip()
MAIL_PASSWORD = os.environ.get("MAIL_PASSWORD", "").strip()
MAIL_USE_TLS = os.environ.get("MAIL_USE_TLS", "true").lower() == "true"
MAIL_FROM = os.environ.get("MAIL_FROM") or MAIL_USERNAME or "no-reply@lab.local"
MAIL_REPLY_TO = os.environ.get("MAIL_REPLY_TO", MAIL_FROM)

logger = logging.getLogger("lab_inventory")
logger.setLevel(logging.INFO)
if not logger.handlers:
    file_handler = logging.FileHandler(LOG_PATH, encoding="utf-8")
    file_handler.setLevel(logging.INFO)
    file_handler.setFormatter(logging.Formatter("%(asctime)s %(levelname)s %(message)s"))
    logger.addHandler(file_handler)


def is_admin_session():
    return session.get("is_admin") is True


def verify_admin_credentials(username, password):
    if not ADMIN_USERNAME or not ADMIN_PASSWORD_HASH:
        return False
    return username == ADMIN_USERNAME and check_password_hash(ADMIN_PASSWORD_HASH, password)


def load_json(path: Path, default):
    if not path.exists():
        path.write_text(json.dumps(default, indent=2), encoding="utf-8")
    with path.open("r", encoding="utf-8") as f:
        try:
            return json.load(f)
        except json.JSONDecodeError:
            path.write_text(json.dumps(default, indent=2), encoding="utf-8")
            return default


def save_json(path: Path, payload):
    path.write_text(json.dumps(payload, indent=2), encoding="utf-8")


def normalized_int(value):
    if value is None:
        return 0
    if isinstance(value, str):
        value = value.strip()
        if not value:
            return 0
        try:
            return int(float(value))
        except ValueError:
            return 0
    return int(value)


def load_inventory():
    workbook = load_workbook(WORKBOOK_PATH, data_only=True)
    sheet = workbook["Inventory"]
    inventory = []

    for row in sheet.iter_rows(min_row=2, values_only=True):
        if row is None or row[0] is None:
            continue
        name = row[0]
        code = row[1]
        available = normalized_int(row[2])
        bin_name = row[3] or "Unassigned"
        closet = row[4] or "Unknown"
        condition = row[5] or "Unknown"
        notes = row[6] or ""
        total = normalized_int(row[7])

        if name is None:
            continue

        inventory.append(
            {
                "item_name": str(name).strip(),
                "code": "" if code is None else str(code).strip(),
                "available_qty": available,
                "bin": str(bin_name).strip(),
                "closet": str(closet).strip(),
                "condition": str(condition).strip(),
                "notes": str(notes).strip(),
                "total_qty": total,
            }
        )

    return inventory


def load_locations():
    workbook = load_workbook(WORKBOOK_PATH, data_only=True)
    locations = {}

    bins_sheet = workbook["Bins"]
    for row in bins_sheet.iter_rows(min_row=2, values_only=True):
        if row and row[0]:
            bin_name = str(row[0]).strip()
            closet = row[1]
            row_num = row[2]
            position = row[3]
            locations[bin_name] = {
                "closet": str(closet).strip() if closet else "Unknown",
                "row": row_num,
                "position": position,
            }

    return locations


def get_status_summary(inventory_items):
    approved_requests = [
        req for req in load_json(REQUESTS_PATH, []) if req.get("status") == "approved"
    ]
    approved_totals = {}
    for request in approved_requests:
        item_name = request.get("item_name")
        if not item_name:
            continue
        approved_totals[item_name] = approved_totals.get(item_name, 0) + normalized_int(request.get("qty"))

    total_items = len(inventory_items)
    total_available = sum(max(0, item["available_qty"] - approved_totals.get(item["item_name"], 0)) for item in inventory_items)
    return total_items, total_available, approved_totals


def make_location_label(item):
    bin_name = item["bin"]
    closet_name = item["closet"]
    if bin_name and closet_name:
        return f"{closet_name} / {bin_name}"
    if closet_name:
        return closet_name
    return "Unassigned"


def get_approved_borrowed_parts():
    approved_requests = [
        req for req in load_json(REQUESTS_PATH, []) if req.get("status") == "approved"
    ]
    borrowed_parts = []
    for req in approved_requests:
        borrowed_parts.append(
            {
                "item_name": req.get("item_name"),
                "qty": req.get("qty"),
                "borrower_name": req.get("borrower_name"),
                "approved_at": req.get("approved_at"),
                "return_date": req.get("return_date"),
                "location": req.get("location"),
            }
        )
    return borrowed_parts


def is_valid_future_date(date_string):
    if not date_string:
        return False
    try:
        return date.fromisoformat(date_string) > date.today()
    except ValueError:
        return False


def build_email_body(subject, message_body, sender_name="Lab Inventory Management"):
    message_body = (message_body or "").strip()
    if not message_body:
        message_body = "This is an automated message from the Lab Inventory Management system."

    plain = f"{message_body}\n\nBest regards,\n{sender_name}"
    html_body = "<html><body style='font-family: Arial, sans-serif; color: #1f2937; line-height: 1.6;'>"
    html_body += f"<h2 style='margin-bottom: 16px; color: #0f172a;'>{subject}</h2>"
    formatted = message_body.replace("\n\n", "</p><p>").replace("\n", "<br>")
    html_body += f"<p>{formatted}</p>"
    html_body += "<p style='margin-top: 18px;'>Best regards,<br><strong>Lab Inventory Management</strong></p>"
    html_body += "</body></html>"
    return plain, html_body


def send_email(to_email, subject, body):
    if not to_email:
        logger.warning("Email send skipped because no recipient was provided. Subject: %s", subject)
        return False
    logger.info("Preparing to send email to %s | Subject: %s", to_email, subject)
    if not MAIL_SERVER or not MAIL_USERNAME or not MAIL_PASSWORD:
        log_message = f"[EMAIL] To: {to_email}\nSubject: {subject}\n{body}\n"
        logger.info(log_message)
        print(log_message)
        return True

    try:
        plain_text, html_text = build_email_body(subject, body)
        message = MIMEMultipart("alternative")
        message["Subject"] = subject
        message["From"] = MAIL_FROM
        message["To"] = to_email
        message["Reply-To"] = MAIL_REPLY_TO
        message["Date"] = formatdate(localtime=True)
        message["X-Mailer"] = "Lab Inventory Manager"
        message["X-Entity-Ref-ID"] = f"lab-inventory-{abs(hash(subject + to_email))}"
        message["Organization"] = "Lab Inventory Management"
        message["Auto-Submitted"] = "auto-generated"

        part1 = MIMEText(plain_text, "plain", "utf-8")
        part2 = MIMEText(html_text, "html", "utf-8")
        message.attach(part1)
        message.attach(part2)

        with smtplib.SMTP(MAIL_SERVER, MAIL_PORT, timeout=20) as server:
            server.ehlo()
            if MAIL_USE_TLS:
                server.starttls()
            server.login(MAIL_USERNAME, MAIL_PASSWORD)
            server.send_message(message)
        logger.info("Email sent successfully to %s | Subject: %s", to_email, subject)
        return True
    except Exception as exc:
        logger.exception("Email failed for %s | Subject: %s", to_email, subject)
        print(f"[EMAIL ERROR] Could not send to {to_email}: {exc}")
        return False


def get_item_approved_total(item_name):
    if not item_name:
        return 0
    return sum(
        normalized_int(req.get("qty"))
        for req in load_json(REQUESTS_PATH, [])
        if req.get("status") == "approved" and (req.get("item_name") or "").lower() == item_name.lower()
    )


@app.get("/logo.svg")
def serve_logo():
    return send_from_directory(ROOT, "industrial-robot-factory-svgrepo-com.svg", mimetype="image/svg+xml")


@app.before_request
def log_request():
    logger.info("REQUEST %s %s", request.method, request.path)


@app.errorhandler(Exception)
def handle_exception(error):
    logger.exception("Unhandled error for %s %s", request.method, request.path)
    if isinstance(error, HTTPException):
        return error
    return "Internal Server Error", 500


@app.get("/")
def index():
    inventory_items = load_inventory()
    location_map = load_locations()
    requests = load_json(REQUESTS_PATH, [])
    approved_requests = [req for req in requests if req.get("status") == "approved"]

    approved_totals = {}
    for request_entry in approved_requests:
        item_name = request_entry.get("item_name")
        if not item_name:
            continue
        approved_totals[item_name] = approved_totals.get(item_name, 0) + normalized_int(request_entry.get("qty"))

    visible_items = []
    for item in inventory_items:
        available_now = max(0, item["available_qty"] - approved_totals.get(item["item_name"], 0))
        location = item["closet"]
        if item["bin"]:
            location = f"{item['closet']} / {item['bin']}"
        details = location_map.get(item["bin"], {})
        visible_items.append(
            {
                **item,
                "available_now": available_now,
                "location": location,
                "row_info": details.get("row"),
                "position": details.get("position"),
            }
        )

    search_query = (request.args.get("search") or "").strip().lower()
    closet_filter = (request.args.get("closet") or "").strip().lower()
    sort_option = request.args.get("sort") or "item-asc"

    filtered_items = visible_items
    if search_query:
        filtered_items = [
            item for item in filtered_items if search_query in (item.get("item_name") or "").lower()
        ]
    if closet_filter:
        filtered_items = [
            item for item in filtered_items if (item.get("closet") or "").lower() == closet_filter
        ]

    if sort_option == "available-desc":
        filtered_items = sorted(filtered_items, key=lambda item: item["available_now"], reverse=True)
    elif sort_option == "closet-asc":
        filtered_items = sorted(filtered_items, key=lambda item: ((item.get("closet") or "").lower(), (item.get("bin") or "").lower()))
    elif sort_option == "bin-asc":
        filtered_items = sorted(filtered_items, key=lambda item: ((item.get("bin") or "").lower(), (item.get("closet") or "").lower()))
    else:
        filtered_items = sorted(filtered_items, key=lambda item: (item.get("item_name") or "").lower())

    total_items, total_available, _ = get_status_summary(inventory_items)
    per_page = 10
    total_pages = max(1, (len(filtered_items) + per_page - 1) // per_page)
    page = max(1, int(request.args.get("page", 1)))
    page = min(page, total_pages)
    start = (page - 1) * per_page
    end = start + per_page
    page_items = filtered_items[start:end]
    page_numbers = list(range(max(1, page - 2), min(total_pages, page + 2) + 1))
    prev_page = max(1, page - 2)
    next_page = min(total_pages, page + 2)

    closet_values = sorted({item.get("closet") for item in visible_items if item.get("closet")})

    return render_template(
        "index.html",
        items=page_items,
        all_items=visible_items,
        total_items=total_items,
        total_available=total_available,
        admin_logged_in=is_admin_session(),
        current_page=page,
        total_pages=total_pages,
        page_size=per_page,
        page_numbers=page_numbers,
        prev_page=prev_page,
        next_page=next_page,
        search_value=search_query,
        selected_closet=closet_filter,
        selected_sort=sort_option,
        closets=closet_values,
    )


@app.post("/borrow")
def submit_borrow_request():
    item_name = (request.form.get("item_name") or "").strip()
    borrower_name = (request.form.get("borrower_name") or "").strip()
    borrower_id = (request.form.get("borrower_id") or "").strip()
    borrower_email = (request.form.get("borrower_email") or "").strip()
    quantity = normalized_int(request.form.get("quantity"))
    purpose = (request.form.get("purpose") or "").strip()
    return_date = (request.form.get("return_date") or "").strip()

    if not item_name or not borrower_name or quantity <= 0:
        return redirect(url_for("index", error="Please enter a valid borrower and quantity."))

    if not is_valid_future_date(return_date):
        return redirect(url_for("index", error="Return date must be in the future."))

    inventory = load_inventory()
    item = next((entry for entry in inventory if entry["item_name"].lower() == item_name.lower()), None)
    if item is None:
        return redirect(url_for("index", error="Selected item was not found."))

    approved_totals = {}
    for request_record in load_json(REQUESTS_PATH, []):
        if request_record.get("status") == "approved" and request_record.get("item_name", "").lower() == item_name.lower():
            approved_totals[item_name] = approved_totals.get(item_name, 0) + normalized_int(request_record.get("qty"))

    available_now = max(0, item["available_qty"] - approved_totals.get(item_name, 0))
    if quantity <= 0 or quantity > available_now:
        return redirect(url_for("index", error="Borrow quantity must be greater than 0 and no more than the available amount."))

    request_list = load_json(REQUESTS_PATH, [])
    new_request = {
        "id": f"req-{len(request_list) + 1}",
        "item_name": item["item_name"],
        "qty": quantity,
        "borrower_name": borrower_name,
        "borrower_id": borrower_id,
        "borrower_email": borrower_email,
        "purpose": purpose,
        "return_date": return_date,
        "location": make_location_label(item),
        "status": "pending",
        "submitted_at": datetime.utcnow().strftime("%Y-%m-%d %H:%M:%S"),
    }
    request_list.append(new_request)
    save_json(REQUESTS_PATH, request_list)
    send_email(
        borrower_email,
        "Borrow Request Submitted",
        f"Dear {borrower_name},\n\nYour request for {quantity} of {item['item_name']} has been submitted and is awaiting admin approval.\nReturn date: {return_date}.",
    )
    return redirect(url_for("index"))


@app.get("/login")
def login_page():
    if is_admin_session():
        return redirect(url_for("admin"))
    return render_template("login.html")


@app.post("/login")
def login_submit():
    username = (request.form.get("username") or "").strip()
    password = request.form.get("password") or ""

    if verify_admin_credentials(username, password):
        session["is_admin"] = True
        return redirect(url_for("admin"))

    return render_template("login.html", error="Invalid username or password."), 401


@app.get("/logout")
def logout():
    session.clear()
    return redirect(url_for("login_page"))


@app.get("/admin")
def admin():
    if not is_admin_session():
        return redirect(url_for("login_page"))

    requests = load_json(REQUESTS_PATH, [])
    pending = [req for req in requests if req.get("status") == "pending"]
    approved = [req for req in requests if req.get("status") == "approved"]
    admin_error = session.pop("admin_error", None)
    return render_template("admin.html", pending=pending, approved=approved, admin_error=admin_error)


@app.post("/admin/cancel/<request_id>")
def cancel_request(request_id):
    if not is_admin_session():
        return redirect(url_for("login_page"))

    requests = load_json(REQUESTS_PATH, [])
    borrower_email = ""
    borrower_name = ""
    for request in requests:
        if request.get("id") == request_id:
            borrower_email = request.get("borrower_email") or ""
            borrower_name = request.get("borrower_name") or "Borrower"
            request["status"] = "cancelled"
            request["cancelled_at"] = datetime.utcnow().strftime("%Y-%m-%d %H:%M:%S")
            break
    save_json(REQUESTS_PATH, requests)
    save_json(BORROWED_PATH, [req for req in requests if req.get("status") == "approved"])
    send_email(
        borrower_email,
        "Borrow Request Cancelled",
        f"Dear {borrower_name},\n\nYour borrow request for {request.get('item_name', 'the selected item')} has been cancelled.",
    )
    return redirect(url_for("admin"))


@app.post("/admin/approve/<request_id>")
def approve_request(request_id):
    if not is_admin_session():
        return redirect(url_for("login_page"))

    requests = load_json(REQUESTS_PATH, [])
    item_name = None
    qty = 0
    borrower_email = ""
    borrower_name = ""
    for request in requests:
        if request.get("id") == request_id:
            item_name = (request.get("item_name") or "").strip()
            qty = normalized_int(request.get("qty"))
            borrower_email = request.get("borrower_email") or ""
            borrower_name = request.get("borrower_name") or "Borrower"
            break

    if item_name:
        item = next((entry for entry in load_inventory() if entry["item_name"].lower() == item_name.lower()), None)
        if item is not None:
            available_now = max(0, item["available_qty"] - get_item_approved_total(item_name))
            if qty > available_now:
                session["admin_error"] = (
                    f"Cannot approve {qty} of {item_name}. Only {available_now} is available to borrow right now."
                )
                return redirect(url_for("admin"))

    for request in requests:
        if request.get("id") == request_id:
            request["status"] = "approved"
            request["approved_at"] = datetime.utcnow().strftime("%Y-%m-%d %H:%M:%S")
            break
    save_json(REQUESTS_PATH, requests)
    save_json(BORROWED_PATH, [req for req in requests if req.get("status") == "approved"])
    send_email(
        borrower_email,
        "Borrow Request Approved",
        f"Dear {borrower_name},\n\nYour request for {qty} of {item_name} has been approved.\nPlease return it by {next((req.get('return_date') for req in requests if req.get('id') == request_id), 'the stated date')}.",
    )
    return redirect(url_for("admin"))


@app.post("/admin/reject/<request_id>")
def reject_request(request_id):
    if not is_admin_session():
        return redirect(url_for("login_page"))

    requests = load_json(REQUESTS_PATH, [])
    borrower_email = ""
    borrower_name = ""
    updated = []
    for request in requests:
        if request.get("id") == request_id:
            borrower_email = request.get("borrower_email") or ""
            borrower_name = request.get("borrower_name") or "Borrower"
            request["status"] = "rejected"
        updated.append(request)
    save_json(REQUESTS_PATH, updated)
    save_json(BORROWED_PATH, [req for req in updated if req.get("status") == "approved"])
    send_email(
        borrower_email,
        "Borrow Request Rejected",
        f"Dear {borrower_name},\n\nYour request for {next((req.get('item_name') for req in updated if req.get('id') == request_id), 'the requested item')} has been rejected by the admin.",
    )
    return redirect(url_for("admin"))


@app.get("/admin/borrowed")
def borrowed_page():
    if not is_admin_session():
        return redirect(url_for("login_page"))

    requests = load_json(REQUESTS_PATH, [])
    approved = [req for req in requests if req.get("status") == "approved"]
    returned = [req for req in requests if req.get("status") == "returned"]
    return render_template("borrowed.html", approved=approved, returned=returned)


@app.post("/admin/return/<request_id>")
def return_request(request_id):
    if not is_admin_session():
        return redirect(url_for("login_page"))

    requests = load_json(REQUESTS_PATH, [])
    for request in requests:
        if request.get("id") == request_id:
            request["status"] = "returned"
            request["returned_at"] = datetime.utcnow().strftime("%Y-%m-%d %H:%M:%S")
            send_email(
                request.get("borrower_email") or "",
                "Borrow Return Confirmed",
                f"Dear {request.get('borrower_name') or 'Borrower'},\n\nThe borrowed item {request.get('item_name')} has been marked as returned.",
            )
            break
    save_json(REQUESTS_PATH, requests)
    save_json(BORROWED_PATH, [req for req in requests if req.get("status") == "approved"])
    return redirect(url_for("borrowed_page"))


if __name__ == "__main__":
    load_json(REQUESTS_PATH, [])
    load_json(BORROWED_PATH, [])
    app.run(debug=True, host="0.0.0.0", port=5000)
