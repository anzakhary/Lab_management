import openpyxl
from pathlib import Path

p = Path('Lab_Inventory_Tracker (2).xlsx')
print('exists', p.exists())
wb = openpyxl.load_workbook(p, data_only=True)
print('sheets', wb.sheetnames)
for sheet in wb.sheetnames:
    ws = wb[sheet]
    print('SHEET', sheet, 'rows', ws.max_row, 'cols', ws.max_column)
    for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 10), values_only=True):
        print(row)
    print('-' * 80)
